"""Pruebas de calidad de datos, captura diaria e informe anual."""

import io
from datetime import date
from unittest import mock

from django.core import mail
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import override_settings
from django.urls import reverse

from investigacion.models import ArticuloCientifico, ArticuloCientificoAutor
from nucleo.fusion import fusionar
from nucleo.models import ConfirmacionInforme, Evidencia, Institucion, PeriodoInforme, Persona, Revista
from nucleo.normalizacion import normalizar
from nucleo.tests import Datos
from vinculacion.models import ArbitrajePublicacion

PREFIJO = 'articulocientificoautor_set'


def datos_articulo(revista, **extra):
    return {'titulo': 'Nuevo', 'revista': revista.pk, 'status': 'ENVIADO', 'fecha_enviado': '10/01/2020',
            f'{PREFIJO}-TOTAL_FORMS': 0, f'{PREFIJO}-INITIAL_FORMS': 0,
            'nucleo-evidencia-content_type-object_id-TOTAL_FORMS': 0,
            'nucleo-evidencia-content_type-object_id-INITIAL_FORMS': 0, **extra}


class DuplicadosTests(Datos):
    def test_aviso_de_persona_parecida_y_confirmacion(self):
        existente = Persona.objects.create(nombre='Juan Carlos', apellidos='Pérez Gómez')
        self.client.force_login(self.ana)
        url = reverse('admin:nucleo_persona_add')
        respuesta = self.client.post(url, {'nombre': 'J. C.', 'apellidos': 'Perez Gomez'})
        self.assertContains(respuesta, 'Ya existen registros parecidos')
        self.assertContains(respuesta, str(existente))
        respuesta = self.client.post(url, {'nombre': 'J. C.', 'apellidos': 'Perez Gomez', 'confirmar_no_duplicado': 'on'})
        self.assertEqual(respuesta.status_code, 302)

    def test_doi_repetido(self):
        otro = self.articulo('Registrado por Beto', self.beto.persona, doi='10.1000/abc')
        self.client.force_login(self.ana)
        respuesta = self.client.post(reverse('admin:investigacion_articulocientifico_add'),
                                     datos_articulo(self.revista, titulo='Otro título', doi='https://doi.org/10.1000/ABC'))
        self.assertContains(respuesta, 'registrado por otro académico')
        self.assertNotContains(respuesta, reverse('admin:investigacion_articulocientifico_change', args=[otro.pk]))

    def test_fusionar_personas(self):
        duplicada = Persona.objects.create(nombre='A.', apellidos='López Pérez')
        articulo = self.articulo('Compartido', self.ana.persona, duplicada)
        solo_duplicada = self.articulo('Solo la duplicada', duplicada)
        fusionar(self.ana.persona, [duplicada])
        self.assertFalse(Persona.objects.filter(pk=duplicada.pk).exists())
        self.assertEqual(list(articulo.autores.all()), [self.ana.persona])
        self.assertEqual(list(solo_duplicada.autores.all()), [self.ana.persona])

    def test_accion_fusionar_solo_para_administradores(self):
        a = Revista.objects.create(nombre='Rev A', pais=self.mexico)
        b = Revista.objects.create(nombre='Rev A.', pais=self.mexico)
        self.articulo('En B', self.ana.persona).__class__.objects.filter(titulo='En B').update(revista=b)
        url = reverse('admin:nucleo_revista_changelist')
        self.client.force_login(self.admin)
        confirmacion = self.client.post(url, {'action': 'fusionar_registros', '_selected_action': [a.pk, b.pk]})
        self.assertContains(confirmacion, 'Elige el registro que se conserva')
        self.client.post(url, {'action': 'fusionar_registros', '_selected_action': [a.pk, b.pk], 'conservar': a.pk})
        self.assertFalse(Revista.objects.filter(pk=b.pk).exists())
        self.assertEqual(ArticuloCientifico.objects.get(titulo='En B').revista, a)
        self.client.force_login(self.ana)
        self.assertNotContains(self.client.get(url), 'fusionar_registros')


class RevisarDuplicadosTests(Datos):
    def test_lista_pares_y_fusiona(self):
        original = Persona.objects.create(nombre='Miguel Ángel', apellidos='Salinas Melgoza')
        copia = Persona.objects.create(nombre='Miguel A.', apellidos='Salinas Melgoza')
        self.articulo('De la copia', copia)
        url = reverse('admin:nucleo_persona_revisar_duplicados')
        self.client.force_login(self.admin)
        self.assertContains(self.client.get(url), 'Miguel A. Salinas Melgoza')
        self.client.post(url, {'conservar': original.pk, 'eliminar': copia.pk})
        self.assertFalse(Persona.objects.filter(pk=copia.pk).exists())
        self.assertEqual(list(ArticuloCientifico.objects.get(titulo='De la copia').autores.all()), [original])
        self.client.force_login(self.ana)
        self.assertEqual(self.client.get(url).status_code, 403)


class ValidacionesTests(Datos):
    def test_fecha_fuera_de_rango(self):
        self.client.force_login(self.ana)
        respuesta = self.client.post(reverse('admin:investigacion_articulocientifico_add'),
                                     datos_articulo(self.revista, fecha_enviado='01/01/2077'))
        self.assertContains(respuesta, 'Revisa el año')

    def test_arbitraje_de_articulo_requiere_revista(self):
        from django.core.exceptions import ValidationError
        arbitraje = ArbitrajePublicacion(tipo='ARTICULO', fecha_dictamen=date(2020, 1, 1), usuario=self.ana)
        with self.assertRaises(ValidationError) as error:
            arbitraje.full_clean()
        self.assertIn('revista', error.exception.message_dict)


class PeriodoCerradoTests(Datos):
    def setUp(self):
        PeriodoInforme.objects.create(anio=2018, fecha_limite=date(2019, 1, 31), cerrado=True)
        self.publicado = self.articulo('Publicado en 2018', self.ana.persona)

    def test_investigador_no_modifica_ni_crea_en_anio_cerrado(self):
        self.client.force_login(self.ana)
        url = reverse('admin:investigacion_articulocientifico_change', args=[self.publicado.pk])
        self.assertFalse(self.client.get(url).context['has_change_permission'])
        respuesta = self.client.post(reverse('admin:investigacion_articulocientifico_add'),
                                     datos_articulo(self.revista, status='PUBLICADO', fecha_publicado='01/06/2018'))
        self.assertContains(respuesta, 'El informe 2018 ya está cerrado')

    def test_administrador_si_puede(self):
        self.client.force_login(self.admin)
        url = reverse('admin:investigacion_articulocientifico_change', args=[self.publicado.pk])
        self.assertTrue(self.client.get(url).context['has_change_permission'])


@override_settings(PRIVATE_MEDIA_ROOT='/tmp/sia-pruebas-privado')
class EvidenciasTests(Datos):
    def test_subir_y_descargar_con_permisos(self):
        articulo = self.articulo('Con constancia', self.ana.persona)
        self.client.force_login(self.ana)
        url = reverse('admin:investigacion_articulocientifico_change', args=[articulo.pk])
        prefijo_evidencia = 'nucleo-evidencia-content_type-object_id'
        autores = list(articulo.articulocientificoautor_set.all())
        respuesta = self.client.post(url, {
            'titulo': articulo.titulo, 'revista': self.revista.pk, 'status': 'PUBLICADO',
            'fecha_publicado': '01/05/2018',
            f'{PREFIJO}-TOTAL_FORMS': 1, f'{PREFIJO}-INITIAL_FORMS': 1,
            f'{PREFIJO}-0-id': autores[0].pk, f'{PREFIJO}-0-articulo': articulo.pk,
            f'{PREFIJO}-0-persona': self.ana.persona.pk, f'{PREFIJO}-0-orden': 0,
            f'{prefijo_evidencia}-TOTAL_FORMS': 1, f'{prefijo_evidencia}-INITIAL_FORMS': 0,
            f'{prefijo_evidencia}-0-archivo': SimpleUploadedFile('constancia.pdf', b'%PDF-1.4 prueba'),
            f'{prefijo_evidencia}-0-descripcion': 'Constancia',
        })
        self.assertEqual(respuesta.status_code, 302, getattr(respuesta, 'context', None) and
                         respuesta.context['adminform'].form.errors)
        evidencia = Evidencia.objects.get()
        self.assertEqual(evidencia.subido_por, self.ana)
        self.assertTrue(evidencia.archivo.url.startswith('/admin/evidencia/'))
        self.assertEqual(b''.join(self.client.get(evidencia.archivo.url).streaming_content), b'%PDF-1.4 prueba')
        self.client.force_login(self.beto)
        self.assertEqual(self.client.get(evidencia.archivo.url).status_code, 403)


class InformeAnualTests(Datos):
    def setUp(self):
        self.periodo = PeriodoInforme.objects.create(anio=2018, fecha_limite=date(2019, 1, 31))
        self.articulo('Artículo 2018', self.ana.persona)

    def test_mi_informe_y_confirmacion(self):
        self.client.force_login(self.ana)
        respuesta = self.client.get(reverse('admin:informe'))
        self.assertContains(respuesta, 'Artículo 2018')
        self.client.post(reverse('admin:informe') + '?anio=2018', {'comentario': 'Completo'})
        self.assertEqual(ConfirmacionInforme.objects.get().comentario, 'Completo')

    def test_avance_y_excel_para_administradores(self):
        self.client.force_login(self.admin)
        avance = self.client.get(reverse('admin:informe_avance'))
        self.assertContains(avance, 'Ana López Pérez')
        excel = self.client.get(reverse('admin:informe_excel'), {'anio': 2018})
        self.assertTrue(excel.content.startswith(b'PK'))
        from openpyxl import load_workbook
        libro = load_workbook(io.BytesIO(excel.content))
        self.assertIn('Investigación', libro.sheetnames)
        self.client.force_login(self.ana)
        self.assertEqual(self.client.get(reverse('admin:informe_avance')).status_code, 403)

    def test_inicio_muestra_pendientes(self):
        ArticuloCientifico.objects.create(titulo='Estancado', revista=self.revista, status='ENVIADO',
                                          fecha_enviado=date(2020, 1, 1))
        ArticuloCientificoAutor.objects.create(articulo=ArticuloCientifico.objects.get(titulo='Estancado'),
                                               persona=self.ana.persona)
        self.client.force_login(self.ana)
        respuesta = self.client.get(reverse('admin:index'))
        self.assertContains(respuesta, 'Estancado')
        self.assertContains(respuesta, 'Informe 2018')


class ImportacionTests(Datos):
    CROSSREF = {'message': {
        'title': ['Soil  erosion in Michoacán'], 'container-title': ['Investigaciones Geográficas'],
        'ISSN': ['01884611'], 'volume': '12', 'issue': '3', 'page': '45-60', 'DOI': '10.1000/xyz',
        'URL': 'https://doi.org/10.1000/xyz', 'published-print': {'date-parts': [[2021, 5]]},
        'author': [{'given': 'Ana', 'family': 'López Pérez'}, {'given': 'Pedro', 'family': 'Nuevo'}],
    }}

    def test_doi_llena_el_formulario(self):
        self.client.force_login(self.ana)
        url = reverse('admin:investigacion_articulocientifico_importar')
        with mock.patch('investigacion.importacion._obtener_json', return_value=self.CROSSREF):
            respuesta = self.client.post(url, {'doi': 'https://doi.org/10.1000/XYZ'})
        self.assertEqual(respuesta.status_code, 302)
        self.assertIn('revista=', respuesta.url)  # La revista se encontró por nombre.
        nuevo = Persona.objects.get(apellidos='Nuevo')
        self.assertFalse(nuevo.verificado)
        alta = self.client.get(respuesta.url)
        formset = next(f for f in alta.context['inline_admin_formsets'] if f.formset.model is ArticuloCientificoAutor)
        self.assertEqual([str(f.initial['persona']) for f in formset.formset.forms[:2]], [str(self.ana.persona.pk), str(nuevo.pk)])
        self.assertEqual(alta.context['adminform'].form.initial['titulo'], 'Soil erosion in Michoacán')

    def test_bibtex(self):
        from investigacion.importacion import desde_bibtex
        [datos] = desde_bibtex('@article{a1, title={Paisajes {de} Michoacán}, author={López, Ana and Pedro Ruiz}, '
                               'journal={Investigaciones Geográficas}, year={2019}, pages={10--20}, doi={10.1/ABC}}')
        self.assertEqual(datos.titulo, 'Paisajes de Michoacán')
        self.assertEqual(datos.autores, [('Ana', 'López'), ('Pedro', 'Ruiz')])
        self.assertEqual((datos.pagina_inicio, datos.pagina_fin, datos.doi), (10, 20, '10.1/abc'))


class HistorialTests(Datos):
    def test_bitacora_protegida(self):
        articulo = self.articulo('Con historia', self.beto.persona)
        articulo.titulo = 'Con historia (v2)'
        articulo.save()
        version = articulo.history.first()
        self.client.force_login(self.beto)
        self.assertEqual(self.client.get(reverse('admin:investigacion_articulocientifico_history',
                                                 args=[articulo.pk])).status_code, 200)
        self.client.force_login(self.ana)
        url = reverse('admin:investigacion_articulocientifico_simple_history', args=[articulo.pk, version.pk])
        self.assertEqual(self.client.get(url).status_code, 403)


class RecuperacionTests(Datos):
    def test_envia_correo_con_enlace(self):
        self.ana.email = 'ana@ejemplo.mx'
        self.ana.save()
        self.client.post(reverse('admin_password_reset'), {'email': 'ana@ejemplo.mx'})
        self.assertEqual(len(mail.outbox), 1)
        self.assertIn('/admin/recuperar/', mail.outbox[0].body)
        self.assertContains(self.client.get(reverse('admin:login')), reverse('admin_password_reset'))


class NormalizacionTests(Datos):
    def test_jerarquia_ambito_y_doi(self):
        from django.apps import apps
        Institucion.objects.create(nombre='Facultad de Ciencias, UNAM', pais=self.mexico)
        self.articulo('Con URL como DOI', self.ana.persona, doi='http://revista.mx/articulo')
        normalizar(apps.get_model, self.mexico.pk)
        facultad = Institucion.objects.get(nombre='Facultad de Ciencias')
        self.assertEqual(facultad.padre, self.institucion)
        articulo = ArticuloCientifico.objects.get(titulo='Con URL como DOI')
        self.assertEqual((articulo.doi, articulo.url), ('', 'http://revista.mx/articulo'))


class ConfiguracionEntidadTests(Datos):
    def test_documentos_usan_la_configuracion(self):
        from nucleo.models import ConfiguracionEntidad, Evento, TipoEvento
        from formatos.models import LicenciaGoceSueldo
        configuracion = ConfiguracionEntidad.actual()
        configuracion.consejo_tecnico = 'Consejo Técnico de Prueba'
        configuracion.save()
        evento = Evento.objects.create(nombre='Congreso', tipo=TipoEvento.objects.create(nombre='Congreso'),
                                       fecha_inicio=date(2026, 5, 1), fecha_fin=date(2026, 5, 3), pais=self.mexico)
        licencia = LicenciaGoceSueldo.objects.create(
            usuario=self.ana, evento=evento, tipo_participacion='Ponente', fecha_inicio=date(2026, 5, 1),
            fecha_fin=date(2026, 5, 3), importancia='Alta', costo='100.00')
        self.client.force_login(self.ana)
        html = self.client.get(reverse('admin:formatos_licenciagocesueldo_descargar_pdf', args=[licencia.pk]),
                               {'formato': 'html'})
        self.assertContains(html, 'Dra. Titular Prueba')
        self.assertContains(html, 'Consejo Técnico de Prueba')

    def test_ambito_segun_pais_sede(self):
        from nucleo.models import Evento, Pais, TipoEvento
        tipo = TipoEvento.objects.create(nombre='Congreso')
        nacional = Evento.objects.create(nombre='A', tipo=tipo, fecha_inicio=date(2020, 1, 1),
                                         fecha_fin=date(2020, 1, 2), pais=self.mexico)
        chile = Pais.objects.create(nombre='Chile', codigo='CL')
        internacional = Evento.objects.create(nombre='B', tipo=tipo, fecha_inicio=date(2020, 1, 1),
                                              fecha_fin=date(2020, 1, 2), pais=chile)
        self.assertEqual((nacional.ambito, internacional.ambito), ('NACIONAL', 'INTERNACIONAL'))

    def test_solo_administradores_y_un_registro(self):
        from nucleo.models import ConfiguracionEntidad
        self.client.force_login(self.admin)
        lista = self.client.get(reverse('admin:nucleo_configuracionentidad_changelist'))
        self.assertRedirects(lista, reverse('admin:nucleo_configuracionentidad_change',
                                            args=[ConfiguracionEntidad.objects.get().pk]))
        self.assertEqual(self.client.get(reverse('admin:nucleo_configuracionentidad_add')).status_code, 403)
        self.client.force_login(self.ana)
        self.assertEqual(self.client.get(reverse('admin:nucleo_configuracionentidad_changelist')).status_code, 403)
