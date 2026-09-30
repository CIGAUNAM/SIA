"""Vistas propias dentro del admin: evidencias, "Mi informe", avance de captura e informe anual en Excel."""

import io
from datetime import date

from django.contrib import admin, messages
from django.contrib.admin.models import LogEntry
from django.core.exceptions import PermissionDenied
from django.db.models import Max
from django.http import FileResponse, Http404, HttpResponse
from django.shortcuts import get_object_or_404, redirect
from django.template.response import TemplateResponse
from django.urls import reverse

from .admin_base import es_administrador
from .cv import secciones_cv
from .models import ConfirmacionInforme, Evidencia, PeriodoInforme, User

TIPOS_ACADEMICOS = [User.Tipo.INVESTIGADOR, User.Tipo.TECNICO, User.Tipo.POSTDOCTORADO]


def evidencia_view(request, nombre):
    """Descarga una evidencia solo si quien la pide puede ver el registro al que pertenece."""
    evidencia = Evidencia.objects.filter(archivo=nombre).select_related('content_type').first()
    if evidencia is None:
        raise Http404
    model_admin = admin.site._registry.get(evidencia.content_type.model_class())
    if model_admin is None or not model_admin.get_queryset(request).filter(pk=evidencia.object_id).exists():
        raise PermissionDenied
    return FileResponse(evidencia.archivo.open('rb'), as_attachment=False,
                        filename=evidencia.archivo.name.rsplit('/', 1)[-1])


def academicos_del_periodo(anio):
    return [u for u in User.objects.filter(tipo__in=TIPOS_ACADEMICOS, is_active=True).order_by('last_name')
            if (u.ingreso_entidad is None or u.ingreso_entidad.year <= anio)
            and (u.egreso_entidad is None or u.egreso_entidad.year >= anio)]


def _periodo_solicitado(request):
    anio = request.GET.get('anio')
    if anio and anio.isdigit():
        return get_object_or_404(PeriodoInforme, anio=int(anio))
    return PeriodoInforme.abierto_actual() or PeriodoInforme.objects.first()


def mi_informe_view(request):
    """Lo capturado por el académico en el año del informe, con enlaces para corregirlo, y la confirmación."""
    administrador = es_administrador(request.user)
    usuario = request.user
    if request.GET.get('usuario') and administrador:
        usuario = get_object_or_404(User, pk=request.GET['usuario'])
    periodo = _periodo_solicitado(request)

    if request.method == 'POST' and periodo is not None and usuario == request.user:
        if periodo.cerrado:
            messages.error(request, f'El informe {periodo.anio} ya está cerrado.')
        else:
            ConfirmacionInforme.objects.update_or_create(
                periodo=periodo, usuario=usuario, defaults={'comentario': request.POST.get('comentario', '')})
            messages.success(request, f'Confirmaste tu informe {periodo.anio}. ¡Gracias!')
        return redirect(f"{reverse('admin:informe')}?anio={periodo.anio}")

    secciones = secciones_cv(usuario, periodo.anio, periodo.anio) if periodo else []
    contexto = {
        **admin.site.each_context(request),
        'title': f'Informe {periodo.anio}' if periodo else 'Mi informe',
        'periodo': periodo,
        'periodos': PeriodoInforme.objects.all(),
        'usuario_informe': usuario,
        'secciones': secciones,
        'total': sum(len(entradas) for _, _, subsecciones in secciones for _, entradas in subsecciones),
        'confirmacion': ConfirmacionInforme.objects.filter(periodo=periodo, usuario=usuario).first()
        if periodo else None,
        'dias_restantes': (periodo.fecha_limite - date.today()).days if periodo and not periodo.cerrado else None,
    }
    if contexto['confirmacion']:
        contexto['mensaje_confirmacion'] = f"Informe confirmado el {contexto['confirmacion'].fecha:%d/%m/%Y}."
    return TemplateResponse(request, 'admin/nucleo/mi_informe.html', contexto)


def _conteos_por_academico(periodo):
    ultima_actividad = dict(LogEntry.objects.values('user').annotate(ultima=Max('action_time'))
                            .values_list('user', 'ultima'))
    confirmaciones = {c.usuario_id: c for c in periodo.confirmaciones.all()}
    filas = []
    for academico in academicos_del_periodo(periodo.anio):
        secciones = secciones_cv(academico, periodo.anio, periodo.anio)
        por_seccion = {clave: sum(len(e) for _, e in subsecciones) for clave, _, subsecciones in secciones}
        filas.append({
            'academico': academico,
            'secciones': secciones,
            'por_seccion': por_seccion,
            'total': sum(por_seccion.values()),
            'confirmacion': confirmaciones.get(academico.pk),
            'ultima_actividad': ultima_actividad.get(academico.pk),
        })
    return filas


def avance_view(request):
    if not es_administrador(request.user):
        raise PermissionDenied
    periodo = _periodo_solicitado(request)
    filas = _conteos_por_academico(periodo) if periodo else []
    contexto = {
        **admin.site.each_context(request),
        'title': f'Avance de captura {periodo.anio}' if periodo else 'Avance de captura',
        'periodo': periodo,
        'periodos': PeriodoInforme.objects.all(),
        'filas': filas,
        'confirmados': sum(1 for f in filas if f['confirmacion']),
        'sin_registros': sum(1 for f in filas if not f['total']),
        'encabezados': ['Académico', 'Tipo', 'Registros en el año', 'Confirmó', 'Última actividad en el sistema'],
        'excel_url': f"{reverse('admin:informe_excel')}?anio={periodo.anio}" if periodo else '',
    }
    return TemplateResponse(request, 'admin/nucleo/avance.html', contexto)


def informe_excel_view(request):
    """Informe anual de la entidad: resumen, avance por académico y una hoja por sección (sin duplicar coautorías)."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter

    from SIA.tablero import construir_tablero

    if not es_administrador(request.user):
        raise PermissionDenied
    periodo = _periodo_solicitado(request)
    if periodo is None:
        messages.error(request, 'No hay periodos de informe registrados.')
        return redirect('admin:index')
    filas = _conteos_por_academico(periodo)
    negritas = Font(bold=True)

    libro = Workbook()
    resumen = libro.active
    resumen.title = 'Resumen'
    resumen.append([f'Informe anual {periodo.anio}'])
    resumen['A1'].font = Font(bold=True, size=14)
    resumen.append([])
    resumen.append(['Indicador', 'Total de la entidad', 'Promedio por académico', 'Máximo'])
    for celda in resumen[3]:
        celda.font = negritas
    for serie in construir_tablero(request.user, hasta=periodo.anio, ver_total=True)['series']:
        resumen.append([serie['titulo'], serie['total'][-1], serie['promedio'][-1], serie['maximo'][-1]])
    resumen.append([])
    resumen.append(['Académicos', len(filas)])
    resumen.append(['Informes confirmados', sum(1 for f in filas if f['confirmacion'])])

    claves = []
    for fila in filas:
        for clave, titulo, _ in fila['secciones']:
            if clave not in [c for c, _ in claves]:
                claves.append((clave, titulo))
    avance = libro.create_sheet('Por académico')
    avance.append(['Académico', 'Tipo', 'Confirmó', 'Total', *[titulo for _, titulo in claves]])
    for celda in avance[1]:
        celda.font = negritas
    for fila in filas:
        academico = fila['academico']
        confirmo = fila['confirmacion'].fecha.strftime('%d/%m/%Y') if fila['confirmacion'] else 'No'
        avance.append([str(academico), academico.get_tipo_display(), confirmo, fila['total'],
                       *[fila['por_seccion'].get(clave, 0) for clave, _ in claves]])

    for clave, titulo in claves:
        hoja = libro.create_sheet(titulo[:31])
        hoja.append(['Apartado', 'Registro', 'Académicos de la entidad'])
        for celda in hoja[1]:
            celda.font = negritas
        registros = {}
        for fila in filas:
            for clave_fila, _, subsecciones in fila['secciones']:
                if clave_fila != clave:
                    continue
                for subtitulo, entradas in subsecciones:
                    for entrada in entradas:
                        llave = (subtitulo, type(entrada.objeto).__name__, getattr(entrada.objeto, 'pk', id(entrada)))
                        registro = registros.setdefault(llave, [subtitulo, entrada.texto, []])
                        registro[2].append(str(fila['academico']))
        for subtitulo, texto, academicos in registros.values():
            hoja.append([subtitulo, texto, ', '.join(academicos)])
        for columna, ancho in ((1, 35), (2, 100), (3, 40)):
            hoja.column_dimensions[get_column_letter(columna)].width = ancho
        for fila_hoja in hoja.iter_rows(min_row=2):
            for celda in fila_hoja:
                celda.alignment = Alignment(wrap_text=True, vertical='top')
    for hoja in (resumen, avance):
        hoja.column_dimensions['A'].width = 45

    salida = io.BytesIO()
    libro.save(salida)
    respuesta = HttpResponse(salida.getvalue(),
                             content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    respuesta["Content-Disposition"] = f'attachment; filename="informe-{periodo.anio}.xlsx"'
    return respuesta


def mi_perfil_view(request):
    """"Mi perfil": la página de la propia cuenta (datos personales, formación y experiencia)."""
    return redirect('admin:nucleo_user_change', request.user.pk)


def orcid_view(request):
    """Lo que ORCID tiene público de una persona, buscada por ORCID o por correo (botón "Buscar en ORCID")."""
    from django.http import JsonResponse

    from .externos import ErrorServicio, datos_orcid, orcid_por_correo
    from .models import Persona

    orcid, correo = request.GET.get('orcid', '').strip(), request.GET.get('email', '').strip()
    try:
        if orcid:
            datos = datos_orcid(orcid)
        elif correo:
            encontrado = orcid_por_correo(correo)
            if not encontrado:
                return JsonResponse({'error': 'No hay un perfil de ORCID con ese correo público.'}, status=404)
            datos = {'orcid': encontrado[0], 'nombre': encontrado[1], 'email': correo.lower()}
        else:
            return JsonResponse({'error': 'Escribe un ORCID o un correo.'}, status=400)
    except ErrorServicio as error:
        return JsonResponse({'error': str(error)}, status=502)
    existente = Persona.objects.filter(orcid=datos['orcid']).exclude(pk=request.GET.get('persona') or None).first()
    if existente is not None:
        datos['existente'] = str(existente)
    return JsonResponse(datos)


def ver_academico_view(request):
    """"Ver por académico": elegir (o dejar de ver) a un académico para acotar las listas de producción."""
    from django.utils.http import url_has_allowed_host_and_scheme

    from .acotar import SESION

    if not es_administrador(request.user):
        raise PermissionDenied
    destino = request.GET.get('next') or request.POST.get('next') or reverse('admin:index')
    if not url_has_allowed_host_and_scheme(destino, allowed_hosts={request.get_host()}):
        destino = reverse('admin:index')
    if request.GET.get('quitar'):
        request.session.pop(SESION, None)
        messages.info(request, 'Ves de nuevo la producción de todos los académicos.')
        return redirect(destino)
    if request.method == 'POST':
        academico = get_object_or_404(User, pk=request.POST.get('usuario'), is_active=True)
        request.session[SESION] = academico.pk
        return redirect(destino)
    academicos = (User.objects.filter(is_active=True, groups__name='Académicos').select_related('persona')
                  .order_by('persona__nombre'))
    contexto = {**admin.site.each_context(request), 'title': 'Ver por académico', 'academicos': academicos,
                'destino': destino}
    return TemplateResponse(request, 'admin/nucleo/ver_academico.html', contexto)
