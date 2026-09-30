"""Currículum vitae del académico en PDF, Word o HTML.

Las secciones se arman en Python como datos estructurados: cada `Entrada` es una lista de
segmentos `(texto, estilo)` con estilo '' (normal), 'b' (negritas) o 'i' (cursivas), más el
rango de años que cubre para poder filtrar por periodo. Las plantillas de salida solo los recorren.
"""

import io
from dataclasses import dataclass
from datetime import date

from django import forms
from django.contrib import admin, messages
from django.core.exceptions import PermissionDenied
from django.db.models import Q
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect
from django.template.response import TemplateResponse
from django.urls import reverse
from django.utils.text import slugify
from unfold.widgets import (UnfoldAdminCheckboxSelectMultipleWidget, UnfoldAdminIntegerFieldWidget,
                            UnfoldAdminRadioSelectWidget, UnfoldAdminSelectWidget)

from .admin_base import es_administrador, persona_de
from .documentos import ErrorDocumento, respuesta_documento
from .models import ConfiguracionEntidad, Libro, LibroParticipante, NivelAcademico, StatusPublicacion, User
from .utils import personas_ordenadas, prefetch_personas


# ---------------------------------------------------------------------------
# Entradas y formato
# ---------------------------------------------------------------------------

@dataclass
class Entrada:
    segmentos: list
    inicio: int | None = None
    fin: int | None = None
    objeto: object = None

    @property
    def url(self):
        """Enlace a la edición del registro en el admin."""
        if self.objeto is None:
            return None
        opts = self.objeto._meta
        return reverse(f'admin:{opts.app_label}_{opts.model_name}_change', args=[self.objeto.pk])

    @property
    def texto(self):
        return ''.join(texto for texto, _ in self.segmentos)

    def en_periodo(self, desde, hasta):
        if self.inicio is None:
            return True
        fin = self.fin if self.fin is not None else date.today().year
        return (hasta is None or self.inicio <= hasta) and (desde is None or fin >= desde)


def b(texto):
    return [(str(texto), 'b')] if texto else []


def i(texto):
    return [(str(texto), 'i')] if texto else []


def _segmentos(parte):
    if not parte:
        return []
    return [(str(parte), '')] if isinstance(parte, str) else list(parte)


def _sin_punto_final(segmentos):
    texto, estilo = segmentos[-1]
    texto = texto.rstrip()
    if texto.endswith('.'):
        texto = texto[:-1]
    return [*segmentos[:-1], (texto, estilo)]


def _unir(*partes):
    """Une las partes con '. ' sin duplicar puntos; cada parte es texto o una lista de segmentos."""
    resultado = []
    for parte in partes:
        segmentos = _segmentos(parte)
        if not ''.join(t for t, _ in segmentos).strip():
            continue
        if resultado:
            resultado.append(('. ', ''))
        resultado.extend(_sin_punto_final(segmentos))
    return [*resultado, ('.', '')] if resultado else []


def _personas(obj, campo='autores'):
    return ', '.join(p.nombre for p in personas_ordenadas(obj, campo))


def _anio(fecha):
    return fecha.year if fecha else None


def _anio_texto(fecha):
    return str(fecha.year) if fecha else 's.f.'


def _texto_periodo(obj):
    if obj.fecha_inicio is None:
        return 's.f.'  # Sin fecha (dato heredado que nadie ha completado).
    inicio = obj.fecha_inicio.year
    fin = obj.fecha_fin.year if obj.fecha_fin else 'actual'
    return f'{inicio}' if fin == inicio else f'{inicio}–{fin}'


def _paginas(obj):
    if obj.pagina_inicio and obj.pagina_fin:
        return f'pp. {obj.pagina_inicio}–{obj.pagina_fin}'
    return ''


def _status(obj):
    return '' if obj.status == StatusPublicacion.PUBLICADO else f'[{obj.get_status_display()}]'


def en_fecha(obj, fecha, *partes):
    return Entrada(_unir(*partes), _anio(fecha), _anio(fecha), obj)


def en_periodo(obj, *partes):
    fin = obj.fecha_fin.year if obj.fecha_fin else None
    return Entrada(_unir(*partes, _texto_periodo(obj)), _anio(obj.fecha_inicio), fin, obj)


def _publicacion(obj, *medio):
    return en_fecha(obj, obj.fecha, f'{_personas(obj)} ({_anio_texto(obj.fecha)})', b(obj.titulo), *medio, _status(obj))


def _articulo(obj):
    volumen = (obj.volumen or '') + (f'({obj.numero})' if obj.numero else '')
    medio = [*i(obj.revista.nombre), *([(f', {volumen}', '')] if volumen else []),
             *([(f', {_paginas(obj)}', '')] if _paginas(obj) else [])]
    doi = f'DOI: {obj.doi}' if getattr(obj, 'doi', '') else ''
    return _publicacion(obj, medio, doi)


def _autores_libro(libro):
    participantes = list(libro.libroparticipante_set.all())
    autores = [p for p in participantes if p.rol == LibroParticipante.Rol.AUTOR] or participantes
    texto = ', '.join(p.persona.nombre for p in autores)
    if autores and autores[0].rol != LibroParticipante.Rol.AUTOR:
        texto += f' ({autores[0].get_rol_display().lower()})'
    return texto


def _libro(libro):
    editorial = ', '.join(x for x in (libro.editorial, libro.ciudad, str(libro.pais)) if x)
    return en_fecha(libro, libro.fecha, f'{_autores_libro(libro)} ({_anio_texto(libro.fecha)})', b(libro.titulo), editorial,
                    f'ISBN {libro.isbn}' if libro.isbn else '', _status(libro))


def _capitulo(obj):
    libro = obj.libro
    return en_fecha(obj, libro.fecha, f'{_personas(obj)} ({_anio_texto(libro.fecha)})', b(obj.titulo),
                    [('En: ', ''), *i(libro.titulo)], libro.editorial, _paginas(obj))


def _con_institucion(parte, obj, campo='institucion'):
    institucion = getattr(obj, campo, None)
    return [*_segmentos(parte), (f'. {institucion.nombre}', '')] if institucion else parte


# ---------------------------------------------------------------------------
# Secciones
# ---------------------------------------------------------------------------

def secciones_cv(usuario, desde=None, hasta=None, incluir=None):
    """Lista de `(clave, título, [(subtítulo, [Entrada])])`, solo con secciones que tienen contenido.

    `desde`/`hasta` filtran por año; `incluir` es un conjunto de claves de sección (None = todas).
    """
    from compromiso_institucional.models import (ApoyoInstitucional, ComisionInstitucional,
                                                 LaborDirectivaCoordinacion, RepresentacionOrganoColegiado)
    from desarrollo_tecnologico.models import DesarrolloTecnologico
    from difusion_cientifica.models import (MemoriaInExtenso, OrganizacionEventoAcademico,
                                            ParticipacionEventoAcademico)
    from distinciones.models import ComisionExpertos, DistincionAcademico, DistincionAlumno, SociedadCientifica
    from divulgacion_cientifica.models import (ArticuloDivulgacion, CapituloLibroDivulgacion,
                                               OrganizacionEventoDivulgacion, ParticipacionEventoDivulgacion,
                                               ProgramaMedio)
    from docencia.models import ArticuloDocencia, CursoEscolarizado, CursoExtracurricular, ProgramaEstudio
    from experiencia_profesional.models import CapacidadPotencialidad, ExperienciaProfesional, LineaInvestigacion
    from formacion_academica.models import CursoEspecializacion, Grado, Postdoctorado
    from formacion_recursos_humanos.models import (AsesoriaEstudiante, ComiteCandidaturaDoctoral, ComiteTutoral,
                                                   DireccionTesis, GrupoInvestigacionInterno,
                                                   SupervisionPostdoctoral)
    from investigacion.models import (ArticuloCientifico, CapituloLibroInvestigacion, MapaArbitrado,
                                      ProyectoInvestigacion, PublicacionTecnica)
    from movilidad_academica.models import MovilidadAcademica
    from vinculacion.models import ArbitrajePublicacion, Convenio, OtraComision, RedAcademica, ServicioAsesoriaExterna

    u = usuario
    p = persona_de(usuario)

    def con_autores(qs, campo='autores'):
        return qs.prefetch_related(prefetch_personas(qs.model, campo)).distinct()

    def libros(tipo):
        qs = Libro.objects.filter(tipo=tipo, participantes=p).distinct().prefetch_related(
            'libroparticipante_set__persona')
        return [_libro(x) for x in qs.con_fecha().order_by('-fecha_orden')]

    orden_grado = {NivelAcademico.DOCTORADO: 0, NivelAcademico.MAESTRIA: 1, NivelAcademico.LICENCIATURA: 2}
    grados = sorted(Grado.objects.filter(usuario=u).select_related('institucion'),
                    key=lambda g: (orden_grado[g.nivel], g.fecha_grado))
    tesis = DireccionTesis.objects.filter(Q(director=p) | Q(codirector=p) | Q(tutores=p)).distinct().select_related(
        'asesorado', 'institucion')

    def entrada_tesis(t):
        anio = _anio_texto(t.fecha_examen) if t.fecha_examen else 'en proceso'
        return Entrada(_unir(f'{t.asesorado} ({anio})', i(t.titulo_tesis), t.get_nivel_display(),
                             t.institucion.nombre),
                       _anio(t.fecha_inicio), _anio(t.fecha_examen), t)

    secciones = [
        ('Formación académica', [
            ('Grados académicos', [
                en_fecha(g, g.fecha_grado, b(g.titulo_obtenido), g.institucion.nombre, str(g.fecha_grado.year),
                         [('Tesis: ', ''), *i(g.titulo_tesis)] if g.titulo_tesis else '', g.distincion_obtenida)
                for g in grados]),
            ('Estancias postdoctorales', [
                en_periodo(x, b(x.titulo_proyecto), x.institucion.nombre)
                for x in Postdoctorado.objects.filter(usuario=u).select_related('institucion')]),
            ('Cursos de especialización', [
                en_periodo(x, f'{x.get_tipo_display()}: {x.nombre} ({x.horas} h)', x.institucion.nombre)
                for x in CursoEspecializacion.objects.filter(usuario=u).select_related('institucion')]),
        ]),
        ('Experiencia profesional', [
            ('Trayectoria', [
                en_periodo(x, _con_institucion(b(x.cargo), x))
                for x in ExperienciaProfesional.objects.filter(usuario=u).select_related('institucion')]),
            ('Líneas de investigación', [
                Entrada(_segmentos(x.nombre), objeto=x) for x in LineaInvestigacion.objects.filter(usuario=u)]),
            ('Capacidades y potencialidades', [
                Entrada(_segmentos(x.nombre), objeto=x) for x in CapacidadPotencialidad.objects.filter(usuario=u)]),
        ]),
        ('Compromiso institucional', [
            ('Labores directivas y de coordinación', [
                en_periodo(x, _con_institucion(str(x.cargo), x))
                for x in LaborDirectivaCoordinacion.objects.filter(usuario=u).select_related('cargo', 'institucion')]),
            ('Representación ante órganos colegiados', [
                en_periodo(x, _con_institucion(str(x), x))
                for x in RepresentacionOrganoColegiado.objects.filter(usuario=u).select_related('institucion')]),
            ('Comisiones institucionales', [
                en_periodo(x, _con_institucion(str(x.comision), x))
                for x in ComisionInstitucional.objects.filter(usuario=u).select_related('comision', 'institucion')]),
            ('Apoyo institucional', [
                en_periodo(x, str(x.actividad), x.descripcion)
                for x in ApoyoInstitucional.objects.filter(usuario=u).select_related('actividad')]),
        ]),
        ('Investigación', [
            ('Proyectos como responsable', [
                en_periodo(x, b(x.nombre), x.get_financiamiento_display(), x.financiamiento_clave)
                for x in ProyectoInvestigacion.objects.filter(responsables=p).distinct()]),
            ('Proyectos como participante', [
                en_periodo(x, b(x.nombre), x.get_financiamiento_display())
                for x in ProyectoInvestigacion.objects.filter(participantes=p).exclude(responsables=p).distinct()]),
            ('Artículos científicos', [
                _articulo(x) for x in con_autores(ArticuloCientifico.objects.filter(autores=p)).select_related(
                    'revista')]),
            ('Libros', libros(Libro.Tipo.INVESTIGACION)),
            ('Capítulos en libros', [
                _capitulo(x) for x in con_autores(CapituloLibroInvestigacion.objects.filter(autores=p)).select_related(
                    'libro')]),
            ('Mapas arbitrados', [
                _publicacion(x, x.publicacion) for x in con_autores(MapaArbitrado.objects.filter(autores=p))]),
            ('Publicaciones técnicas', [
                _publicacion(x, x.get_tipo_display()) for x in con_autores(
                    PublicacionTecnica.objects.filter(autores=p))]),
            ('Desarrollos tecnológicos', [
                en_fecha(x, x.fecha, f'{_personas(x)} ({_anio_texto(x.fecha)})', b(x.nombre),
                         f'Versión {x.version}' if x.version else '', x.licencia)
                for x in con_autores(DesarrolloTecnologico.objects.filter(autores=p))]),
        ]),
        ('Difusión científica', [
            ('Memorias in extenso', [
                en_fecha(x, x.fecha, f'{_personas(x)} ({_anio_texto(x.fecha)})', b(x.titulo), i(x.evento),
                         f'{x.ciudad}, {x.pais}' if x.ciudad else str(x.pais), _paginas(x))
                for x in con_autores(MemoriaInExtenso.objects.filter(autores=p)).select_related('pais')]),
            ('Participación en eventos académicos', [
                en_fecha(x, x.fecha, f'{_personas(x)} ({_anio_texto(x.fecha)})', b(x.titulo),
                         [(f'{x.get_tipo_display()} en ', ''), *i(x.evento)], str(x.pais),
                         'Por invitación' if x.por_invitacion else '',
                         'Conferencia magistral' if x.ponencia_magistral else '')
                for x in con_autores(ParticipacionEventoAcademico.objects.filter(autores=p)).select_related('pais')]),
            ('Organización de eventos académicos', [
                en_fecha(x, x.evento.fecha_inicio, b(x.evento.nombre), x.get_tipo_participacion_display(),
                         _anio_texto(x.evento.fecha_inicio))
                for x in OrganizacionEventoAcademico.objects.filter(usuario=u).select_related('evento')]),
        ]),
        ('Divulgación científica', [
            ('Artículos de divulgación', [
                _articulo(x) for x in con_autores(ArticuloDivulgacion.objects.filter(autores=p)).select_related(
                    'revista')]),
            ('Libros', libros(Libro.Tipo.DIVULGACION)),
            ('Capítulos en libros', [
                _capitulo(x) for x in con_autores(CapituloLibroDivulgacion.objects.filter(autores=p)).select_related(
                    'libro')]),
            ('Participación en eventos de divulgación', [
                en_fecha(x, x.fecha, f'{_personas(x)} ({_anio_texto(x.fecha)})', b(x.titulo), i(x.evento.nombre))
                for x in con_autores(ParticipacionEventoDivulgacion.objects.filter(autores=p)).select_related(
                    'evento')]),
            ('Organización de eventos de divulgación', [
                en_fecha(x, x.evento.fecha_inicio, b(x.evento.nombre), x.get_tipo_participacion_display(),
                         _anio_texto(x.evento.fecha_inicio))
                for x in OrganizacionEventoDivulgacion.objects.filter(usuario=u).select_related('evento')]),
            ('Medios de comunicación', [
                en_fecha(x, x.fecha, b(x.tema), f'{x.get_actividad_display()} en {x.medio}', _anio_texto(x.fecha))
                for x in ProgramaMedio.objects.filter(usuario=u).select_related('medio')]),
        ]),
        ('Vinculación', [
            ('Arbitraje de publicaciones', [
                en_fecha(x, x.fecha_dictamen, x.get_tipo_display(), x.revista.nombre if x.revista else x.obra,
                         _anio_texto(x.fecha_dictamen))
                for x in ArbitrajePublicacion.objects.filter(usuario=u).select_related('revista')]),
            ('Otras comisiones de arbitraje', [
                en_periodo(x, _con_institucion(str(x), x))
                for x in OtraComision.objects.filter(usuario=u).select_related('tipo', 'institucion')]),
            ('Redes académicas', [
                Entrada(_unir(b(x.nombre), x.get_ambito_display(), _anio_texto(x.fecha_constitucion)),
                        _anio(x.fecha_constitucion), _anio(x.fecha_fin), x)
                for x in RedAcademica.objects.filter(participantes=p)]),
            ('Convenios', [en_periodo(x, b(x.nombre)) for x in Convenio.objects.filter(participantes=p)]),
            ('Servicios y asesorías externas', [
                en_periodo(x, _con_institucion(b(x.nombre), x))
                for x in ServicioAsesoriaExterna.objects.filter(usuario=u).select_related('institucion')]),
        ]),
        ('Movilidad académica', [
            (tipo.label, [
                en_periodo(x, _con_institucion(x.academico, x))
                for x in MovilidadAcademica.objects.filter(usuario=u, tipo=tipo).select_related('institucion')])
            for tipo in MovilidadAcademica.Tipo
        ]),
        ('Docencia', [
            ('Cursos escolarizados', [
                en_periodo(x, b(str(x.asignatura)), x.programa.nombre if x.programa else x.get_nivel_display(),
                           x.institucion.nombre, f'{x.periodo_academico} ({x.total_horas} h)',
                           x.get_nombramiento_display())
                for x in CursoEscolarizado.objects.filter(usuario=u).select_related(
                    'asignatura', 'programa', 'institucion')]),
            ('Cursos extracurriculares', [
                en_periodo(x, f'{x.get_tipo_display()}: {x.asignatura} ({x.total_horas} h)', x.institucion.nombre)
                for x in CursoExtracurricular.objects.filter(usuario=u).select_related('asignatura', 'institucion')]),
            ('Artículos para docencia', [
                _publicacion(x) for x in con_autores(ArticuloDocencia.objects.filter(autores=p))]),
            ('Libros para docencia', libros(Libro.Tipo.DOCENCIA)),
            ('Programas de estudio', [
                en_fecha(x, x.fecha, b(x.nombre), x.get_nivel_display(), x.institucion.nombre, _anio_texto(x.fecha))
                for x in ProgramaEstudio.objects.filter(usuario=u).select_related('institucion')]),
        ]),
        ('Formación de recursos humanos', [
            ('Tesis concluidas', [entrada_tesis(t) for t in tesis if t.status == DireccionTesis.Status.TERMINADA]),
            ('Tesis en proceso', [entrada_tesis(t) for t in tesis if t.status != DireccionTesis.Status.TERMINADA]),
            ('Asesorías de estudiantes', [
                en_periodo(x, f'{x.asesorado}: {x.get_tipo_display()}', x.get_nivel_display())
                for x in AsesoriaEstudiante.objects.filter(usuario=u).select_related('asesorado')]),
            ('Supervisión de investigadores postdoctorales', [
                en_periodo(x, str(x.investigador), i(x.titulo_proyecto))
                for x in SupervisionPostdoctoral.objects.filter(usuario=u).select_related('investigador')]),
            ('Comités tutorales', [
                en_periodo(x, str(x.estudiante), x.get_nivel_display())
                for x in ComiteTutoral.objects.filter(miembros=p).distinct().select_related('estudiante')]),
            ('Comités de candidatura doctoral', [
                en_fecha(x, x.fecha_defensa, str(x.candidato), _anio_texto(x.fecha_defensa))
                for x in ComiteCandidaturaDoctoral.objects.filter(
                    Q(director=p) | Q(codirector=p) | Q(asesores=p) | Q(miembros=p)).distinct().select_related(
                    'candidato')]),
            ('Grupos de investigación internos', [
                en_periodo(x, x.nombre) for x in GrupoInvestigacionInterno.objects.filter(integrantes=p)]),
        ]),
        ('Distinciones', [
            ('Distinciones recibidas', [
                en_fecha(x, x.fecha, b(str(x.distincion)),
                         x.distincion.institucion.nombre if x.distincion.institucion else '', _anio_texto(x.fecha))
                for x in DistincionAcademico.objects.filter(usuario=u).select_related('distincion__institucion')]),
            ('Distinciones de alumnos tutorados', [
                en_fecha(x, x.fecha, f'{x.alumno}: {x.distincion}', _anio_texto(x.fecha))
                for x in DistincionAlumno.objects.filter(tutores=p).select_related('alumno', 'distincion')]),
            ('Comisiones de expertos', [
                en_periodo(x, _con_institucion(x.nombre, x))
                for x in ComisionExpertos.objects.filter(usuario=u).select_related('institucion')]),
            ('Sociedades científicas', [
                en_periodo(x, x.nombre, x.get_tipo_display()) for x in SociedadCientifica.objects.filter(usuario=u)]),
        ]),
    ]

    resultado = []
    for titulo, subsecciones in secciones:
        clave = slugify(titulo)
        if incluir is not None and clave not in incluir:
            continue
        subsecciones = [(subtitulo, [e for e in entradas if e.en_periodo(desde, hasta)])
                        for subtitulo, entradas in subsecciones]
        subsecciones = [(subtitulo, entradas) for subtitulo, entradas in subsecciones if entradas]
        if subsecciones:
            resultado.append((clave, titulo, subsecciones))
    return resultado


TITULOS_SECCIONES = [
    'Formación académica', 'Experiencia profesional', 'Compromiso institucional', 'Investigación',
    'Difusión científica', 'Divulgación científica', 'Vinculación', 'Movilidad académica', 'Docencia',
    'Formación de recursos humanos', 'Distinciones',
]


def datos_generales(usuario):
    return [(etiqueta, valor) for etiqueta, valor in (
        ('Correo electrónico', usuario.email),
        ('Nivel SNII', usuario.get_sni_display()),
        ('Nivel PRIDE', usuario.pride),
        ('Ingreso a la UNAM', usuario.ingreso_unam and usuario.ingreso_unam.strftime('%d/%m/%Y')),
        ('Página web', usuario.url),
    ) if valor]


# ---------------------------------------------------------------------------
# Salidas
# ---------------------------------------------------------------------------

def cv_docx(usuario, secciones, subtitulo, entidad):
    from docx import Document
    from docx.shared import Pt

    documento = Document()
    documento.styles['Normal'].font.name = 'Calibri'
    documento.styles['Normal'].font.size = Pt(10.5)
    documento.add_heading(f'{usuario.grado} {usuario}'.strip(), level=0)
    documento.add_paragraph(f'{entidad.nombre_completo} · {subtitulo}')
    for etiqueta, valor in datos_generales(usuario):
        parrafo = documento.add_paragraph()
        parrafo.add_run(f'{etiqueta}: ').bold = True
        parrafo.add_run(str(valor))
    if usuario.semblanza:
        documento.add_heading('Semblanza', level=2)
        documento.add_paragraph(usuario.semblanza)
    for numero, (_, titulo, subsecciones) in enumerate(secciones, start=1):
        documento.add_heading(f'{numero}. {titulo}', level=1)
        for subnumero, (titulo_subseccion, entradas) in enumerate(subsecciones, start=1):
            documento.add_heading(f'{numero}.{subnumero}. {titulo_subseccion}', level=2)
            for indice, entrada in enumerate(entradas, start=1):
                parrafo = documento.add_paragraph()
                parrafo.paragraph_format.left_indent = Pt(18)
                parrafo.paragraph_format.first_line_indent = Pt(-18)
                parrafo.add_run(f'{indice}.\t')
                for texto, estilo in entrada.segmentos:
                    run = parrafo.add_run(texto)
                    run.bold = estilo == 'b'
                    run.italic = estilo == 'i'
    salida = io.BytesIO()
    documento.save(salida)
    return salida.getvalue()


class OpcionesCV(forms.Form):
    FORMATOS = [('pdf', 'PDF'), ('docx', 'Word (.docx)'), ('html', 'Vista previa en el navegador')]

    usuario = forms.ModelChoiceField(
        queryset=User.objects.filter(is_active=True).order_by('first_name', 'last_name'), required=False,
        widget=UnfoldAdminSelectWidget, label='Académico')
    desde = forms.IntegerField(required=False, min_value=1950, max_value=2100, widget=UnfoldAdminIntegerFieldWidget,
                               label='Desde el año', help_text='Vacío: desde el inicio.')
    hasta = forms.IntegerField(required=False, min_value=1950, max_value=2100, widget=UnfoldAdminIntegerFieldWidget,
                               label='Hasta el año', help_text='Vacío: hasta hoy.')
    secciones = forms.MultipleChoiceField(
        choices=[(slugify(t), t) for t in TITULOS_SECCIONES], required=False,
        widget=UnfoldAdminCheckboxSelectMultipleWidget, label='Secciones', help_text='Sin marcar: todas.')
    formato = forms.ChoiceField(choices=FORMATOS, initial='pdf', widget=UnfoldAdminRadioSelectWidget)

    def clean(self):
        datos = super().clean()
        if datos.get('desde') and datos.get('hasta') and datos['desde'] > datos['hasta']:
            raise forms.ValidationError('El año inicial no puede ser mayor que el final.')
        return datos


def cv_view(request, usuario_id=None):
    """Formulario de opciones del CV; con `generar` en la URL devuelve el documento."""
    administrador = es_administrador(request.user)
    if usuario_id is not None and usuario_id != request.user.pk and not administrador:
        raise PermissionDenied
    form = OpcionesCV(request.GET if 'generar' in request.GET else None,
                      initial={'usuario': usuario_id or request.user.pk, 'formato': 'pdf'})
    if not administrador:
        del form.fields['usuario']

    if form.is_bound and form.is_valid():
        datos = form.cleaned_data
        usuario = datos.get('usuario') or get_object_or_404(User, pk=usuario_id or request.user.pk)
        desde, hasta = datos['desde'], datos['hasta']
        secciones = secciones_cv(usuario, desde, hasta, set(datos['secciones']) or None)
        subtitulo = f"Periodo {desde or 'inicio'}–{hasta or 'actual'}" if desde or hasta else 'Currículum vitae'
        nombre_archivo = f'cv-{slugify(usuario.get_full_name()) or usuario.pk}'
        if datos['formato'] == 'docx':
            respuesta = HttpResponse(
                cv_docx(usuario, secciones, subtitulo, ConfiguracionEntidad.actual(request)),
                content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document')
            respuesta['Content-Disposition'] = f'attachment; filename="{nombre_archivo}.docx"'
            return respuesta
        contexto = {'usuario': usuario, 'datos': datos_generales(usuario), 'secciones': secciones,
                    'subtitulo': subtitulo, 'entidad': ConfiguracionEntidad.actual(request), 'hoy': date.today()}
        try:
            return respuesta_documento(request, 'nucleo/cv.html', contexto, nombre_archivo, datos['formato'])
        except ErrorDocumento as error:
            messages.error(request, str(error))
            return redirect('admin:cv')

    contexto = {**admin.site.each_context(request), 'title': 'Mi currículum', 'form': form}
    return TemplateResponse(request, 'admin/nucleo/cv_opciones.html', contexto)
