import base64
import io
import json
import re

from django import forms
from django.contrib import admin, messages
from django.core.exceptions import PermissionDenied
from django.http import HttpResponse, JsonResponse
from django.shortcuts import get_object_or_404, redirect
from django.template.response import TemplateResponse
from django.urls import path, reverse
from django.utils.html import format_html
from unfold.admin import TabularInline
from unfold.decorators import action
from unfold.widgets import UnfoldAdminTextInputWidget, UnfoldBooleanWidget

from nucleo.admin_base import CatalogoAdmin
from nucleo.models import ConfiguracionEntidad

from .indicadores import calcular
from .models import Emision, Grafica, Informe, validar_periodo


def textos(informe, entidad):
    """Valores para {entidad}, {periodo}… en títulos y pies."""
    inicio, fin = informe.periodo.split('-')
    return {'entidad': entidad.siglas or entidad.nombre, 'institucion': entidad.institucion_madre_siglas or '',
            'periodo': informe.periodo.replace('-', '–'), 'anio': fin, 'inicio': inicio, 'fin': fin}


def rellenar(texto, valores):
    return re.sub(r'\{(\w+)\}', lambda m: str(valores.get(m.group(1), m.group(0))), texto or '')


def datos_informe(informe, entidad):
    """Las gráficas visibles del informe con sus textos ya rellenados y sus datos calculados."""
    valores = textos(informe, entidad)
    graficas = []
    for g in informe.graficas.filter(visible=True):
        try:
            datos = calcular(g, informe.periodo)
        except Exception as error:  # Un indicador que falla no debe tumbar todo el informe.
            datos = {'error': f'No se pudo calcular: {error}'}
        graficas.append({'id': g.pk, 'seccion': g.seccion, 'tipo': g.tipo, 'indicador': g.indicador,
                         'titulo': rellenar(g.titulo, valores), 'subtitulo': rellenar(g.subtitulo, valores),
                         'nota': rellenar(g.nota, valores), 'colores': g.paleta(), 'mostrar_total': g.mostrar_total,
                         'mostrar_valores': g.mostrar_valores, 'apilado': g.apilado, 'datos': datos})
    return {'nombre': informe.nombre, 'periodo': informe.periodo, 'pie': rellenar(informe.pie, valores),
            'graficas': graficas}


class GraficaInline(TabularInline):
    model = Grafica
    fields = ['orden', 'seccion', 'indicador', 'tipo', 'titulo', 'periodos', 'visible']
    ordering_field = 'orden'
    hide_ordering_field = True
    extra = 0
    show_change_link = True


class EmisionInline(TabularInline):
    """Versiones emitidas: solo se consultan (se emiten con el botón «Emitir»)."""
    model = Emision
    fields = ['version', 'emitido_en', 'emitido_por', 'motivo', 'ver_version']
    readonly_fields = fields
    extra = 0
    verbose_name_plural = 'versiones emitidas'

    def has_add_permission(self, request, obj=None):
        return False

    def has_change_permission(self, request, obj=None):
        return False

    def has_delete_permission(self, request, obj=None):
        return False

    @admin.display(description='')
    def ver_version(self, obj):
        return format_html('<a class="text-primary-600" href="{}?version={}">Ver</a> · <a class="text-primary-600" '
                           'href="{}?version={}">Cambios desde entonces</a>',
                           reverse('admin:informes_ver', args=[obj.informe_id]), obj.version,
                           reverse('admin:informes_cambios', args=[obj.informe_id]), obj.version)


class EmitirForm(forms.Form):
    motivo = forms.CharField(required=False, widget=forms.Textarea(attrs={'rows': 4, 'class': (
        'border border-base-200 bg-white rounded-default shadow-xs px-3 py-2 w-full dark:bg-base-900 '
        'dark:border-base-700')}), help_text='Obligatorio si ya hay una versión emitida: qué cambió y por qué.')

    def __init__(self, *args, requiere_motivo=False, **kwargs):
        super().__init__(*args, **kwargs)
        self.requiere_motivo = requiere_motivo

    def clean_motivo(self):
        motivo = self.cleaned_data['motivo'].strip()
        if self.requiere_motivo and not motivo:
            raise forms.ValidationError('Explica por qué se emite una nueva versión.')
        return motivo


class NuevoDesdeForm(forms.Form):
    nombre = forms.CharField(max_length=255, widget=UnfoldAdminTextInputWidget)
    periodo = forms.CharField(max_length=9, validators=[validar_periodo], help_text='Por ejemplo «2026-2027».',
                              widget=UnfoldAdminTextInputWidget)
    es_plantilla = forms.BooleanField(required=False, label='Guardarlo como plantilla', widget=UnfoldBooleanWidget)


@admin.register(Informe)
class InformeAdmin(CatalogoAdmin):
    permisos_investigador = ()
    list_display = ['nombre', 'periodo', 'es_plantilla', 'num_graficas', 'ver']
    list_filter = ['es_plantilla', 'periodo']
    search_fields = ['nombre', 'descripcion']
    fields = ['nombre', 'periodo', 'es_plantilla', 'descripcion', 'pie']
    inlines = [GraficaInline, EmisionInline]
    actions_detail = ['ver_informe', 'emitir', 'nuevo_desde']

    def get_queryset(self, request):
        from django.db.models import Count

        return super().get_queryset(request).annotate(n_graficas=Count('graficas'))

    @admin.display(description='gráficas', ordering='n_graficas')
    def num_graficas(self, obj):
        return obj.n_graficas

    @admin.display(description='')
    def ver(self, obj):
        return format_html('<a href="{}">Ver informe</a>', reverse('admin:informes_ver', args=[obj.pk]))

    def save_model(self, request, obj, form, change):
        if not change:
            obj.creado_por = request.user
        super().save_model(request, obj, form, change)

    @action(description='Ver informe', icon='bar_chart', url_path='ver-informe')
    def ver_informe(self, request, object_id):
        return redirect('admin:informes_ver', object_id)

    @action(description='Emitir', icon='verified', url_path='emitir')
    def emitir(self, request, object_id):
        return redirect('admin:informes_emitir', object_id)

    @action(description='Nuevo informe a partir de este', icon='content_copy', url_path='nuevo-desde')
    def nuevo_desde(self, request, object_id):
        return redirect('admin:informes_nuevo_desde', object_id)

    def get_urls(self):
        vista = self.admin_site.admin_view
        return [
            path('<int:pk>/ver/', vista(self.ver_view), name='informes_ver'),
            path('<int:pk>/nuevo/', vista(self.nuevo_desde_view), name='informes_nuevo_desde'),
            path('<int:pk>/excel/', vista(self.excel_view), name='informes_excel'),
            path('<int:pk>/documento/', vista(self.documento_view), name='informes_documento'),
            path('<int:pk>/emitir/', vista(self.emitir_view), name='informes_emitir'),
            path('<int:pk>/cambios/', vista(self.cambios_view), name='informes_cambios'),
            *super().get_urls(),
        ]

    def _informe(self, request, pk, cambiar=False):
        informe = get_object_or_404(Informe, pk=pk)
        permitido = self.has_change_permission(request, informe) if cambiar else self.has_view_permission(request, informe)
        if not permitido:
            raise PermissionDenied
        return informe

    def _datos(self, request, informe, parametros=None):
        """Datos a mostrar: los de una versión emitida (por omisión, la última) o los actuales (`vivo`)."""
        parametros = parametros if parametros is not None else request.GET
        entidad = ConfiguracionEntidad.actual(request)
        emision = None
        if not parametros.get('vivo'):
            emisiones = informe.emisiones.all()
            emision = (emisiones.filter(version=parametros.get('version')).first() if parametros.get('version')
                       else emisiones.first())
        return (emision.datos if emision else datos_informe(informe, entidad)), emision, entidad

    def ver_view(self, request, pk):
        informe = self._informe(request, pk)
        datos, emision, entidad = self._datos(request, informe)
        diferencias = cambios = 0
        if emision and emision == informe.emisiones.first():
            diferencias = len(diferencias_con_actual(emision, datos_informe(informe, entidad)))
            cambios = len(cambios_desde(emision.emitido_en, limite=500, periodo=informe.periodo))
        return TemplateResponse(request, 'admin/informes/ver.html', {
            **self.admin_site.each_context(request), 'title': informe.nombre, 'opts': self.model._meta,
            'original': informe, 'informe': informe, 'datos': datos, 'entidad': entidad, 'emision': emision,
            'emisiones': informe.emisiones.all(), 'diferencias': diferencias, 'cambios': cambios,
            'parametros': {'version': emision.version} if emision else {'vivo': 1},
            'excel_url': reverse('admin:informes_excel', args=[informe.pk]) +
                         (f'?version={emision.version}' if emision else '?vivo=1'),
            'puede_editar': self.has_change_permission(request, informe), 'version_js': _version_js()})

    def emitir_view(self, request, pk):
        informe = self._informe(request, pk, cambiar=True)
        if informe.es_plantilla:
            messages.error(request, 'Una plantilla no se emite: crea un informe a partir de ella.')
            return redirect('admin:informes_informe_change', informe.pk)
        anterior = informe.emisiones.first()
        form = EmitirForm(request.POST or None, requiere_motivo=anterior is not None)
        if request.method == 'POST' and form.is_valid():
            entidad = ConfiguracionEntidad.actual(request)
            emision = Emision.objects.create(
                informe=informe, version=(anterior.version + 1) if anterior else 1, periodo=informe.periodo,
                emitido_por=request.user, motivo=form.cleaned_data['motivo'], datos=datos_informe(informe, entidad))
            messages.success(request, f'Se emitió la versión {emision.version} de «{informe}». Sus cifras ya no '
                                      'cambian aunque se capturen o corrijan registros.')
            return redirect(f"{reverse('admin:informes_ver', args=[informe.pk])}?version={emision.version}")
        contexto = {**self.admin_site.each_context(request), 'opts': self.model._meta, 'original': informe,
                    'title': f'Emitir «{informe}»', 'form': form, 'informe': informe, 'anterior': anterior,
                    'cambios': len(cambios_desde(anterior.emitido_en, limite=500, periodo=informe.periodo)) if anterior else None}
        return TemplateResponse(request, 'admin/informes/emitir.html', contexto)

    def cambios_view(self, request, pk):
        informe = self._informe(request, pk)
        emision = informe.emisiones.filter(version=request.GET.get('version')).first() or informe.emisiones.first()
        if emision is None:
            messages.info(request, 'Este informe todavía no se ha emitido.')
            return redirect('admin:informes_ver', informe.pk)
        actuales = datos_informe(informe, ConfiguracionEntidad.actual(request))
        contexto = {**self.admin_site.each_context(request), 'opts': self.model._meta, 'original': informe,
                    'title': f'Cambios desde la versión {emision.version}', 'informe': informe, 'emision': emision,
                    'diferencias': diferencias_con_actual(emision, actuales),
                    'cambios': cambios_desde(emision.emitido_en, periodo=informe.periodo),
                    'encabezados': ['Fecha', 'Tipo', 'Qué', 'Registro', 'Quién', 'Motivo'],
                    'url_version': f"{reverse('admin:informes_ver', args=[informe.pk])}?version={emision.version}",
                    'url_vivo': f"{reverse('admin:informes_ver', args=[informe.pk])}?vivo=1"}
        return TemplateResponse(request, 'admin/informes/cambios.html', contexto)

    def nuevo_desde_view(self, request, pk):
        origen = self._informe(request, pk)
        if not self.has_add_permission(request):
            raise PermissionDenied
        if request.method == 'POST':
            form = NuevoDesdeForm(request.POST)
            if form.is_valid():
                copia = origen.duplicar(request.user, **form.cleaned_data)
                messages.success(request, f'Se creó «{copia}» a partir de «{origen}».')
                return redirect('admin:informes_informe_change', copia.pk)
        else:
            inicio = int(origen.periodo[:4]) + (1 if not origen.es_plantilla else 0)
            form = NuevoDesdeForm(initial={'nombre': re.sub(r'\s*\(plantilla\)', '', origen.nombre, flags=re.I),
                                           'periodo': f'{inicio}-{inicio + 1}'})
        contexto = {**self.admin_site.each_context(request), 'title': f'Nuevo informe a partir de «{origen}»',
                    'form': form, 'origen': origen, 'opts': self.model._meta}
        return TemplateResponse(request, 'admin/informes/nuevo_desde.html', contexto)

    def excel_view(self, request, pk):
        from openpyxl import Workbook
        from openpyxl.styles import Font

        informe = self._informe(request, pk)
        datos, emision, _ = self._datos(request, informe)
        libro = Workbook()
        libro.remove(libro.active)
        usados = set()
        for i, g in enumerate(datos['graficas'], start=1):
            nombre = re.sub(r'[\[\]:*?/\\]', '', f'{i}. {g["titulo"]}')[:31]
            while nombre in usados:
                nombre = nombre[:28] + f' {len(usados)}'
            usados.add(nombre)
            hoja = libro.create_sheet(nombre)
            hoja.append([g['titulo']])
            hoja['A1'].font = Font(bold=True, size=13)
            if g['subtitulo']:
                hoja.append([g['subtitulo']])
            for fila in tabla_de(g['datos']):
                hoja.append(fila)
            hoja.column_dimensions['A'].width = 45
        salida = io.BytesIO()
        libro.save(salida)
        respuesta = HttpResponse(salida.getvalue(), content_type='application/vnd.openxmlformats-officedocument.'
                                                                   'spreadsheetml.sheet')
        respuesta['Content-Disposition'] = f'attachment; filename="{_archivo(informe, emision)}.xlsx"'
        return respuesta

    def documento_view(self, request, pk):
        """PDF o Word del informe con las imágenes de las gráficas que dibujó el navegador."""
        informe = self._informe(request, pk)
        if request.method != 'POST':
            return JsonResponse({'error': 'Usa POST'}, status=405)
        cuerpo = json.loads(request.body or '{}')
        datos, emision, entidad = self._datos(request, informe, {k: cuerpo.get(k) for k in ('version', 'vivo')})
        imagenes = {int(k): v for k, v in (cuerpo.get('imagenes') or {}).items()}
        formato = cuerpo.get('formato', 'pdf')
        for g in datos['graficas']:
            g['imagen'] = imagenes.get(g['id'], '')
            g['tabla'] = tabla_de(g['datos'])
        if formato == 'docx':
            contenido, tipo = documento_word(informe, datos, entidad), \
                'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
        else:
            from django.template.loader import render_to_string
            from nucleo.documentos import ErrorDocumento, html_a_pdf

            html = render_to_string('informes/documento.html', {'informe': informe, 'datos': datos,
                                                                'entidad': entidad}, request=request)
            try:
                contenido, tipo = html_a_pdf(html), 'application/pdf'
            except ErrorDocumento as error:
                return JsonResponse({'error': str(error)}, status=500)
        respuesta = HttpResponse(contenido, content_type=tipo)
        respuesta['Content-Disposition'] = f'attachment; filename="{_archivo(informe, emision)}.{formato}"'
        return respuesta


def _version_js():
    """Fecha del script de las gráficas, para que el navegador no use una versión vieja guardada."""
    from pathlib import Path

    archivo = Path(__file__).parent / 'static' / 'informes' / 'graficas.js'
    return int(archivo.stat().st_mtime) if archivo.exists() else 0


def _archivo(informe, emision=None):
    version = f' v{emision.version}' if emision else ' borrador'
    return re.sub(r'[^\w\-]+', '_', f'{informe.nombre} {informe.periodo}{version}').strip('_')


def diferencias_con_actual(emision, actuales):
    """Gráficas cuyas cifras ya no son las emitidas: [(título, [(fila emitida, fila actual), ...])]."""
    vivas = {g['id']: g for g in actuales['graficas']}
    resultado = []
    for g in emision.datos['graficas']:
        actual = vivas.get(g['id'])
        if actual is None or json.dumps(g['datos'], sort_keys=True) == json.dumps(actual['datos'], sort_keys=True):
            continue
        antes, ahora = tabla_de(g['datos']), tabla_de(actual['datos'])
        filas = [(a, b) for a, b in zip(antes + [[]] * (len(ahora) - len(antes)), ahora + [[]] * (len(antes) - len(ahora)))
                 if a != b]
        resultado.append((g['titulo'], filas))
    return resultado


def cambios_desde(momento, limite=300, periodo=None):
    """Registros de producción creados, modificados o borrados después de `momento`, según el historial de cada
    modelo; con `periodo` («2025-2026»), solo los que caen en sus años (o no tienen fecha que los ubique)."""
    from django.apps import apps

    from nucleo.admin_base import PropietarioAdmin
    from nucleo.informe import anio_cierre

    anios = {int(a) for a in periodo.split('-')} if periodo else None

    tipos = {'+': 'Alta', '~': 'Cambio', '-': 'Baja'}
    cambios = []
    for modelo in apps.get_models():
        original = getattr(modelo, 'instance_type', None)
        if not modelo.__name__.startswith('Historical') or original is None or \
                original._meta.app_label in ('informes', 'auth', 'sessions', 'admin', 'contenttypes'):
            continue
        model_admin = admin.site._registry.get(original)
        if not isinstance(model_admin, PropietarioAdmin):
            continue  # Catálogos (personas, revistas…): no son producción del periodo.
        for h in modelo.objects.filter(history_date__gt=momento).select_related('history_user').order_by(
                '-history_date')[:limite]:
            if anios is not None:
                anio = anio_cierre(h.instance, model_admin.campo_fecha)
                if anio is not None and anio not in anios:
                    continue
            enlace = ''
            if h.history_type != '-':
                try:
                    enlace = reverse(f'admin:{original._meta.app_label}_{original._meta.model_name}_change',
                                     args=[h.id])
                except Exception:
                    enlace = ''
            cambios.append({'fecha': h.history_date, 'tipo': tipos.get(h.history_type, h.history_type),
                            'modelo': original._meta.verbose_name, 'registro': str(h.instance)[:120],
                            'usuario': h.history_user, 'motivo': h.history_change_reason or '', 'enlace': enlace})
    cambios.sort(key=lambda c: c['fecha'], reverse=True)
    return cambios[:limite]


def tabla_de(datos):
    """Las cifras de una gráfica como filas (para el Excel y las tablas de los documentos)."""
    if datos.get('error'):
        return [[datos['error']]]
    filas = []
    if datos.get('grupos'):
        filas.append(['Grupo', 'Concepto', 'Valor'])
        for g in datos['grupos']:
            for item in g['items']:
                filas.append([g['nombre'], item['nombre'], item['valor']])
        if datos.get('total') is not None:
            filas.append(['Total', '', datos['total']])
        return filas
    if datos.get('jerarquia'):
        filas.append(['Categoría', 'Subcategoría', 'Detalle', 'Valor'])
        for a in datos['jerarquia']:
            for b in a.get('hijos') or [{'nombre': '', 'valor': a['valor']}]:
                for c in b.get('hijos') or [{'nombre': '', 'valor': b['valor']}]:
                    filas.append([a['nombre'], b['nombre'], c['nombre'], c['valor']])
        if datos.get('total') is not None:
            filas.append(['Total', '', '', datos['total']])
        return filas
    for panel in datos.get('paneles', []):
        if panel.get('titulo'):
            filas.append([panel['titulo'] + (f' (total: {panel["total"]})' if panel.get('total') is not None else '')])
        filas.append(['', *[s['nombre'] for s in panel['series']]])
        for i, categoria in enumerate(panel['categorias']):
            filas.append([categoria, *[s['valores'][i] for s in panel['series']]])
        filas.append([])
    if datos.get('total') is not None:
        filas.append(['Total', datos['total']])
    return filas


def documento_word(informe, datos, entidad):
    from docx import Document
    from docx.shared import Cm, Pt

    doc = Document()
    doc.styles['Normal'].font.name = 'Calibri'
    doc.styles['Normal'].font.size = Pt(10.5)
    doc.add_heading(informe.nombre, 0)
    doc.add_paragraph(f'{entidad.nombre} · Periodo {datos["periodo"].replace("-", "–")}')
    seccion = None
    for g in datos['graficas']:
        if g['seccion'] and g['seccion'] != seccion:
            seccion = g['seccion']
            doc.add_heading(seccion, 1)
        imagen = g.get('imagen') or ''
        if imagen.startswith('data:image/png;base64,'):
            doc.add_picture(io.BytesIO(base64.b64decode(imagen.split(',', 1)[1])), width=Cm(16.5))
        else:
            doc.add_heading(g['titulo'], 2)
            if g['subtitulo']:
                doc.add_paragraph(g['subtitulo'])
            filas = [f for f in g['tabla'] if f]
            ancho = max(len(f) for f in filas) if filas else 0
            if ancho:
                tabla = doc.add_table(rows=0, cols=ancho)
                tabla.style = 'Light Grid Accent 1'
                for fila in filas:
                    celdas = tabla.add_row().cells
                    for i, valor in enumerate(fila):
                        celdas[i].text = '' if valor is None else str(valor)
        if g['nota']:
            doc.add_paragraph(g['nota'])
    salida = io.BytesIO()
    doc.save(salida)
    return salida.getvalue()


@admin.register(Grafica)
class GraficaAdmin(CatalogoAdmin):
    permisos_investigador = ()
    list_display = ['titulo', 'informe', 'indicador', 'tipo', 'orden', 'visible']
    list_filter = ['informe', 'indicador', 'tipo']
    search_fields = ['titulo', 'informe__nombre']
    fieldsets = (
        (None, {'fields': ('informe', 'orden', 'seccion', 'indicador', 'tipo', 'periodos', 'visible')}),
        ('Textos', {'fields': ('titulo', 'subtitulo', 'nota')}),
        ('Estilo', {'fields': ('colores', 'mostrar_total', 'mostrar_valores', 'apilado')}),
    )

    def get_model_perms(self, request):
        return {}  # Se editan desde su informe; no van en el menú.

    def response_change(self, request, obj):
        if '_continue' not in request.POST and '_addanother' not in request.POST:
            return redirect('admin:informes_informe_change', obj.informe_id)
        return super().response_change(request, obj)
